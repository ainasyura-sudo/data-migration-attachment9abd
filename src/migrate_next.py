import os
import getpass
from pathlib import Path

import psycopg
import yaml
from openpyxl import load_workbook


# ============================================================
# 1. PROJECT PATHS / CONFIGURATION
# ============================================================

BASE = Path(__file__).resolve().parent.parent

CONFIG_PATH = BASE / "environment.yaml"

if not CONFIG_PATH.exists():
    raise FileNotFoundError(
        f"Configuration file not found: {CONFIG_PATH}"
    )


with open(CONFIG_PATH, "r", encoding="utf-8") as f:
    config = yaml.safe_load(f) or {}


# ============================================================
# 2. DATABASE CONFIGURATION
# ============================================================

db_cfg = config.get("database", {})

db_host = db_cfg.get("host", "localhost")
db_port = int(db_cfg.get("port", 5432))
db_name = db_cfg.get("name", "postgres")
db_user = db_cfg.get("user", "postgres")
db_schema = db_cfg.get("schema", "public")


# Prefer password from environment.yaml.
# If unavailable, try environment variable.
# Otherwise ask user.
db_password = db_cfg.get("password")

password_env_var = db_cfg.get(
    "password_env",
    "DIC_DB_PASSWORD",
)

if not db_password:
    db_password = os.environ.get(password_env_var)

if not db_password:
    db_password = getpass.getpass(
        "PostgreSQL password: "
    )

if not db_password:
    raise RuntimeError(
        "Database password was not provided."
    )


# ============================================================
# 3. SOURCE WORKBOOK
# ============================================================

source_cfg = config.get("source", {})

excel_file = BASE / source_cfg.get(
    "workbook",
    "AWPB 2026 - Balingian PSC - Attachment 9 2.xlsx",
)

if not excel_file.exists():
    raise FileNotFoundError(
        f"Workbook not found: {excel_file}"
    )


# ============================================================
# 4. TABLE CONFIGURATION
# ============================================================

tables_cfg = config.get("tables", {})

tbl_masterplan = tables_cfg.get(
    "masterplan",
    "abd_masterplan_combined_stg",
)

tbl_decom_summary = tables_cfg.get(
    "decom_summary",
    "abd_decom_summary_stg",
)

tbl_decom_facility = tables_cfg.get(
    "decom_facility",
    "abd_decom_facility_stg",
)

tbl_decom_combined = tables_cfg.get(
    "decom_combined",
    "abd_decom_combined_stg",
)


# ============================================================
# 5. PAC NAME MAPPING CONFIGURATION
# ============================================================

mapping_cfg = config.get("mapping", {})

pac_mapping_file = BASE / mapping_cfg.get(
    "pac_names",
    "pac_name_mapping.yaml",
)

pac_aliases = {}

if pac_mapping_file.exists():

    with open(
        pac_mapping_file,
        "r",
        encoding="utf-8",
    ) as f:

        pac_data = yaml.safe_load(f) or {}

        pac_aliases = pac_data.get(
            "aliases",
            {},
        )


# Workbook does not contain a verified PAC/operator source.
# Do not guess from field name.
operator = None

# Template/source does not currently provide verified values
# for these metadata fields.
sheet_id = None
submission_date = None
region = None


print(
    "PAC name mapping report: "
    "Workbook does not contain a verified PAC name "
    "source column. Operator left NULL."
)


# ============================================================
# 6. HELPER FUNCTIONS
# ============================================================

def number(value):
    """
    Convert Excel numeric value to float.

    None, empty string and "-" become NULL.
    """

    if value is None:
        return None

    if (
        isinstance(value, str)
        and value.strip() in {"", "-"}
    ):
        return None

    return float(value)


# ============================================================
# 7. LOAD WORKBOOK
# ============================================================

wb = load_workbook(
    excel_file,
    data_only=True,
)

ws_9a = wb["Attachment 9a"]
ws_9b = wb["Attachment 9b"]
ws_9g = wb["Attachment 9g"]
ws_9h = wb["Attachment 9h"]


# ============================================================
# 8. ATTACHMENT 9G
#    DECOMMISSIONING COST SUMMARY
# ============================================================

decom_summary = []


for row in range(7, 14):

    field = ws_9g.cell(
        row,
        2,
    ).value

    if not field:
        continue

    decom_summary.append(
        (
            str(field).strip(),

            *[
                number(
                    ws_9g.cell(
                        row,
                        col,
                    ).value
                )
                for col in range(3, 11)
            ],
        )
    )


# ============================================================
# 9. ATTACHMENT 9H
#    DECOMMISSIONING FACILITY BREAKDOWN
# ============================================================

decom_facilities = []

current_field = None


for row in range(
    1,
    ws_9h.max_row + 1,
):

    label = str(
        ws_9h.cell(
            row,
            2,
        ).value
        or ""
    ).strip().lower()


    # --------------------------------------------------------
    # Detect current field
    # --------------------------------------------------------

    if label == "field":

        field_value = ws_9h.cell(
            row,
            3,
        ).value

        current_field = (
            str(field_value).strip()
            if field_value
            else None
        )


    # --------------------------------------------------------
    # Detect facility section
    # --------------------------------------------------------

    elif (
        label == "facility"
        and current_field
    ):

        facility = ws_9h.cell(
            row,
            3,
        ).value

        decom_facilities.append(
            (
                current_field,

                # Keep source row because facility names
                # such as Terminal may repeat.
                row,

                (
                    str(facility).strip()
                    if facility
                    else None
                ),

                number(
                    ws_9h.cell(
                        row,
                        6,
                    ).value
                ),
            )
        )


# ============================================================
# 10. ATTACHMENT 9A(ii)
#     MASTER ABANDONMENT PLAN
# ============================================================

annual_work = {}


for row in range(29, 36):

    field = ws_9a.cell(
        row,
        2,
    ).value

    if not field:
        continue

    field = str(field).strip()


    # Years 2026-2030
    for col in range(3, 8):

        year = ws_9a.cell(
            28,
            col,
        ).value

        if (
            isinstance(
                year,
                (int, float),
            )
            and not isinstance(
                year,
                bool,
            )
        ):

            annual_work[
                (
                    field,
                    int(year),
                )
            ] = number(
                ws_9a.cell(
                    row,
                    col,
                ).value
            )


# ============================================================
# 11. ATTACHMENT 9B
#     ABANDONMENT ACTIVITIES
# ============================================================

activities = []

current_field = None


for row in range(
    6,
    ws_9b.max_row + 1,
):

    field_cell = ws_9b.cell(
        row,
        2,
    ).value


    if field_cell:

        current_field = str(
            field_cell
        ).strip()


    if not current_field:
        continue


    activity = ws_9b.cell(
        row,
        5,
    ).value


    if not activity:
        continue


    asset = ws_9b.cell(
        row,
        3,
    ).value


    remarks = ws_9b.cell(
        row,
        31,
    ).value


    # --------------------------------------------------------
    # Each year contains:
    #
    # Q1
    # Q2
    # Q3
    # Q4
    # FY
    #
    # 2026 -> col 6
    # 2027 -> col 11
    # 2028 -> col 16
    # 2029 -> col 21
    # 2030 -> col 26
    # --------------------------------------------------------

    for year_index, start_col in enumerate(
        (
            6,
            11,
            16,
            21,
            26,
        )
    ):

        year = 2026 + year_index


        quarters = [

            number(
                ws_9b.cell(
                    row,
                    start_col + offset,
                ).value
            )

            for offset in range(4)
        ]


        fy_amount = number(
            ws_9b.cell(
                row,
                start_col + 4,
            ).value
        )


        # Only retain nonzero FY activity.
        if (
            fy_amount is None
            or fy_amount == 0
        ):
            continue


        activities.append(
            (
                current_field,
                year,
                row,

                (
                    str(asset).strip()
                    if asset
                    else None
                ),

                str(activity).strip(),

                *quarters,

                fy_amount,

                (
                    str(remarks).strip()
                    if remarks
                    else None
                ),
            )
        )


# ============================================================
# 12. COMBINE ATTACHMENT 9A(ii) + 9B
# ============================================================

activity_by_key = {}


for activity in activities:

    key = (
        activity[0],
        activity[1],
    )


    if key in activity_by_key:

        raise ValueError(
            "More than one 9b activity for "
            f"{key}. "
            "The masterplan needs an "
            "activity-level key."
        )


    activity_by_key[key] = activity


master_plan = []


for (
    field,
    year,
), work_amount in sorted(
    annual_work.items()
):

    activity = activity_by_key.get(
        (
            field,
            year,
        )
    )


    if activity:

        master_plan.append(
            (
                # field_id
                # resolved from am_field later
                None,

                sheet_id,
                submission_date,
                region,
                operator,

                field,
                year,

                work_amount,

                # source_row_9b onwards
                *activity[2:],
            )
        )


    else:

        master_plan.append(
            (
                # field_id
                None,

                sheet_id,
                submission_date,
                region,
                operator,

                field,
                year,

                work_amount,

                # source_row_9b
                None,

                # asset_9b
                None,

                # activity_9b
                None,

                # Q1
                None,

                # Q2
                None,

                # Q3
                None,

                # Q4
                None,

                # FY
                None,

                # remarks
                None,
            )
        )


# ============================================================
# 13. EXTRACTION SUMMARY
# ============================================================

print(
    f"9g field summaries: "
    f"{len(decom_summary)}"
)

print(
    f"9h facility sections: "
    f"{len(decom_facilities)}"
)

print(
    f"9a(ii) annual plan rows: "
    f"{len(annual_work)}"
)

print(
    f"9b nonzero activity/year rows: "
    f"{len(activities)}"
)

print(
    f"Combined masterplan rows: "
    f"{len(master_plan)}"
)


# ============================================================
# 14. DATABASE CONNECTION
# ============================================================

with psycopg.connect(
    host=db_host,
    port=db_port,
    dbname=db_name,
    user=db_user,
    password=db_password,
    connect_timeout=10,
) as connection:

    with connection.cursor() as cursor:


        # ====================================================
        # 15. RESOLVE field_id FROM am_field
        # ====================================================

        source_fields = sorted(

            {
                row[0]
                for row in decom_summary
            }

            |

            {
                row[5]
                for row in master_plan
            }
        )


        cursor.execute(
            f"""
            SELECT
                field_name,
                id
            FROM {db_schema}.am_field
            WHERE field_name = ANY(%s)
            ORDER BY field_name
            """,

            (
                source_fields,
            ),
        )


        hierarchy_rows = (
            cursor.fetchall()
        )


        field_ids = {}

        duplicate_fields = set()


        for (
            field_name,
            field_id,
        ) in hierarchy_rows:

            if field_name in field_ids:

                duplicate_fields.add(
                    field_name
                )

            else:

                field_ids[
                    field_name
                ] = field_id


        # ----------------------------------------------------
        # Prevent ambiguous hierarchy mapping
        # ----------------------------------------------------

        if duplicate_fields:

            raise RuntimeError(
                "Duplicate exact field names "
                "found in am_field; "
                "mapping is ambiguous: "
                f"{sorted(duplicate_fields)}"
            )


        # ----------------------------------------------------
        # Insert verified field_id into masterplan rows
        #
        # Unregistered fields stay NULL.
        # ----------------------------------------------------

        master_plan = [

            (
                field_ids.get(
                    row[5]
                ),

                *row[1:],
            )

            for row in master_plan
        ]


        mapped = sorted(
            field_ids
        )


        unmapped = sorted(
            set(source_fields)
            - set(field_ids)
        )


        print(
            "Asset hierarchy mapped fields: "
            + (
                ", ".join(mapped)
                or "none"
            )
        )


        print(
            "Fields not in am_field -> "
            "field_id NULL: "
            + (
                ", ".join(unmapped)
                or "none"
            )
        )


        # ====================================================
        # 16. SUPPORTING DECOM TABLES
        #
        # IMPORTANT:
        #
        # These two supporting tables still use "field".
        #
        # Do NOT rename these to field_name here.
        # ====================================================


        # ----------------------------------------------------
        # Attachment 9g summary
        # ----------------------------------------------------

        cursor.execute(
            f"""
            CREATE TABLE IF NOT EXISTS
            {db_schema}.{tbl_decom_summary}
            (
                field TEXT PRIMARY KEY,

                platform_rm_million
                    DOUBLE PRECISION,

                well_rm_million
                    DOUBLE PRECISION,

                pipeline_rm_million
                    DOUBLE PRECISION,

                floater_rm_million
                    DOUBLE PRECISION,

                subsea_facility_rm_million
                    DOUBLE PRECISION,

                terminal_rm_million
                    DOUBLE PRECISION,

                total_rt_rm_million
                    DOUBLE PRECISION,

                total_mod_rm_million
                    DOUBLE PRECISION
            )
            """
        )


        # ----------------------------------------------------
        # Attachment 9h facility
        # ----------------------------------------------------

        cursor.execute(
            f"""
            CREATE TABLE IF NOT EXISTS
            {db_schema}.{tbl_decom_facility}
            (
                field TEXT NOT NULL,

                source_row_9h
                    INTEGER NOT NULL,

                facility TEXT,

                total_cost_rm_million
                    DOUBLE PRECISION,

                PRIMARY KEY (
                    field,
                    source_row_9h
                )
            )
            """
        )


        # ====================================================
        # IMPORTANT
        #
        # We DO NOT recreate:
        #
        # abd_masterplan_combined_stg
        # abd_decom_combined_stg
        #
        # here.
        #
        # Their final UUID / metadata schema already exists
        # in PostgreSQL and is managed separately.
        # ====================================================


        # ====================================================
        # 17. DECOM COMBINED UNIQUE INDEX
        # ====================================================

        cursor.execute(
            f"""
            CREATE UNIQUE INDEX IF NOT EXISTS
                {tbl_decom_combined}_uk
            ON
                {db_schema}.{tbl_decom_combined}
                (
                    field_name,
                    source_row_9h
                )
            """
        )


        # ====================================================
        # 18. UPSERT 9G SUMMARY
        # ====================================================

        cursor.executemany(
            f"""
            INSERT INTO
                {db_schema}.{tbl_decom_summary}
            (
                field,

                platform_rm_million,
                well_rm_million,
                pipeline_rm_million,
                floater_rm_million,

                subsea_facility_rm_million,

                terminal_rm_million,

                total_rt_rm_million,
                total_mod_rm_million
            )

            VALUES (
                %s,
                %s,
                %s,
                %s,
                %s,
                %s,
                %s,
                %s,
                %s
            )

            ON CONFLICT (field)

            DO UPDATE SET

                platform_rm_million =
                    EXCLUDED.platform_rm_million,

                well_rm_million =
                    EXCLUDED.well_rm_million,

                pipeline_rm_million =
                    EXCLUDED.pipeline_rm_million,

                floater_rm_million =
                    EXCLUDED.floater_rm_million,

                subsea_facility_rm_million =
                    EXCLUDED.subsea_facility_rm_million,

                terminal_rm_million =
                    EXCLUDED.terminal_rm_million,

                total_rt_rm_million =
                    EXCLUDED.total_rt_rm_million,

                total_mod_rm_million =
                    EXCLUDED.total_mod_rm_million
            """,

            decom_summary,
        )


        # ====================================================
        # 19. UPSERT 9H FACILITY
        # ====================================================

        cursor.executemany(
            f"""
            INSERT INTO
                {db_schema}.{tbl_decom_facility}
            (
                field,
                source_row_9h,
                facility,
                total_cost_rm_million
            )

            VALUES (
                %s,
                %s,
                %s,
                %s
            )

            ON CONFLICT (
                field,
                source_row_9h
            )

            DO UPDATE SET

                facility =
                    EXCLUDED.facility,

                total_cost_rm_million =
                    EXCLUDED.total_cost_rm_million
            """,

            decom_facilities,
        )


        # ====================================================
        # 20. UPSERT MASTERPLAN COMBINED
        #
        # Final table uses:
        #
        # field_name
        # UUID field_id
        # UUID sheet_id
        #
        # Existing id and created_at are preserved.
        # ====================================================

        cursor.executemany(
            f"""
            INSERT INTO
                {db_schema}.{tbl_masterplan}
            (
                field_id,
                sheet_id,
                submission_date,
                region,
                operator,

                field_name,
                year,

                work_amount_9a_rm_million,

                source_row_9b,
                asset_9b,
                activity_9b,

                q1_rm_million,
                q2_rm_million,
                q3_rm_million,
                q4_rm_million,

                activity_fy_rm_million,

                remarks_9b
            )

            VALUES (
                %s::uuid,
                %s::uuid,
                %s,
                %s,
                %s,
                %s,
                %s,
                %s,
                %s,
                %s,
                %s,
                %s,
                %s,
                %s,
                %s,
                %s,
                %s
            )

            ON CONFLICT (
                field_name,
                year
            )

            DO UPDATE SET

                field_id =
                    EXCLUDED.field_id,

                sheet_id =
                    EXCLUDED.sheet_id,

                submission_date =
                    EXCLUDED.submission_date,

                region =
                    EXCLUDED.region,

                operator =
                    EXCLUDED.operator,

                work_amount_9a_rm_million =
                    EXCLUDED.work_amount_9a_rm_million,

                source_row_9b =
                    EXCLUDED.source_row_9b,

                asset_9b =
                    EXCLUDED.asset_9b,

                activity_9b =
                    EXCLUDED.activity_9b,

                q1_rm_million =
                    EXCLUDED.q1_rm_million,

                q2_rm_million =
                    EXCLUDED.q2_rm_million,

                q3_rm_million =
                    EXCLUDED.q3_rm_million,

                q4_rm_million =
                    EXCLUDED.q4_rm_million,

                activity_fy_rm_million =
                    EXCLUDED.activity_fy_rm_million,

                remarks_9b =
                    EXCLUDED.remarks_9b,

                updated_at =
                    CURRENT_TIMESTAMP
            """,

            master_plan,
        )


        # ====================================================
        # 21. UPSERT DECOM COMBINED
        #
        # Supporting tables:
        #     g.field
        #     h.field
        #
        # Final combined table:
        #     field_name
        # ====================================================

        cursor.execute(
            f"""
            INSERT INTO
                {db_schema}.{tbl_decom_combined}
            (
                field_id,
                sheet_id,
                submission_date,
                region,
                operator,

                field_name,

                platform_rm_million,
                well_rm_million,
                pipeline_rm_million,
                floater_rm_million,

                subsea_facility_rm_million,

                terminal_rm_million,

                total_rt_rm_million,
                total_mod_rm_million,

                source_row_9h,

                facility,

                facility_cost_9h_rm_million
            )

            SELECT

                NULL::uuid
                    AS field_id,

                NULL::uuid
                    AS sheet_id,

                NULL
                    AS submission_date,

                NULL
                    AS region,

                NULL
                    AS operator,

                g.field,

                g.platform_rm_million,
                g.well_rm_million,
                g.pipeline_rm_million,
                g.floater_rm_million,

                g.subsea_facility_rm_million,

                g.terminal_rm_million,

                g.total_rt_rm_million,
                g.total_mod_rm_million,

                COALESCE(
                    h.source_row_9h,
                    0
                )
                    AS source_row_9h,

                h.facility,

                h.total_cost_rm_million
                    AS facility_cost_9h_rm_million

            FROM
                {db_schema}.{tbl_decom_summary}
                AS g

            LEFT JOIN
                {db_schema}.{tbl_decom_facility}
                AS h

                ON g.field = h.field


            ON CONFLICT (
                field_name,
                source_row_9h
            )

            DO UPDATE SET

                field_id =
                    EXCLUDED.field_id,

                sheet_id =
                    EXCLUDED.sheet_id,

                submission_date =
                    EXCLUDED.submission_date,

                region =
                    EXCLUDED.region,

                operator =
                    EXCLUDED.operator,

                platform_rm_million =
                    EXCLUDED.platform_rm_million,

                well_rm_million =
                    EXCLUDED.well_rm_million,

                pipeline_rm_million =
                    EXCLUDED.pipeline_rm_million,

                floater_rm_million =
                    EXCLUDED.floater_rm_million,

                subsea_facility_rm_million =
                    EXCLUDED.subsea_facility_rm_million,

                terminal_rm_million =
                    EXCLUDED.terminal_rm_million,

                total_rt_rm_million =
                    EXCLUDED.total_rt_rm_million,

                total_mod_rm_million =
                    EXCLUDED.total_mod_rm_million,

                facility =
                    EXCLUDED.facility,

                facility_cost_9h_rm_million =
                    EXCLUDED.facility_cost_9h_rm_million,

                updated_at =
                    CURRENT_TIMESTAMP
            """
        )


        # ====================================================
        # 22. APPLY field_id TO DECOM COMBINED
        # ====================================================

        # First force all values to NULL.
        #
        # This ensures fields not registered in am_field
        # cannot retain an old/incorrect field_id.

        cursor.execute(
            f"""
            UPDATE
                {db_schema}.{tbl_decom_combined}

            SET
                field_id = NULL
            """
        )


        # Only populate IDs for exact field-name matches
        # returned from am_field.

        for (
            field_name,
            field_id,
        ) in field_ids.items():

            cursor.execute(
                f"""
                UPDATE
                    {db_schema}.{tbl_decom_combined}

                SET
                    field_id = %s::uuid

                WHERE
                    field_name = %s
                """,

                (
                    str(field_id),
                    field_name,
                ),
            )


        # ====================================================
        # 23. FINAL field_id SAFETY CHECK
        # ====================================================

        for check_table in (
            tbl_masterplan,
            tbl_decom_combined,
        ):


            # ------------------------------------------------
            # Check that every non-NULL field_id exists
            # in am_field.
            # ------------------------------------------------

            cursor.execute(
                f"""
                SELECT DISTINCT

                    s.field_name,
                    s.field_id

                FROM
                    {db_schema}.{check_table}
                    AS s

                LEFT JOIN
                    {db_schema}.am_field
                    AS f

                    ON s.field_id = f.id

                WHERE
                    s.field_id IS NOT NULL

                    AND f.id IS NULL
                """
            )


            invalid = (
                cursor.fetchall()
            )


            if invalid:

                raise RuntimeError(
                    "Invalid field_id detected "
                    f"in {check_table}: "
                    f"{invalid}"
                )


            # ------------------------------------------------
            # If a field does not exist in am_field,
            # field_id MUST remain NULL.
            # ------------------------------------------------

            cursor.execute(
                f"""
                SELECT DISTINCT

                    s.field_name,
                    s.field_id

                FROM
                    {db_schema}.{check_table}
                    AS s

                WHERE NOT EXISTS
                (
                    SELECT 1

                    FROM
                        {db_schema}.am_field
                        AS f

                    WHERE
                        f.field_name =
                        s.field_name
                )

                AND
                    s.field_id IS NOT NULL
                """
            )


            should_be_null = (
                cursor.fetchall()
            )


            if should_be_null:

                raise RuntimeError(
                    "Fields absent from am_field "
                    "must have NULL field_id in "
                    f"{check_table}: "
                    f"{should_be_null}"
                )


        # ====================================================
        # 24. ROW COUNT CONFIRMATION
        # ====================================================

        for tbl_name in (

            tbl_decom_summary,

            tbl_decom_facility,

            tbl_masterplan,

            tbl_decom_combined,

        ):

            cursor.execute(
                f"""
                SELECT COUNT(*)
                FROM {db_schema}.{tbl_name}
                """
            )


            row_count = (
                cursor.fetchone()[0]
            )


            print(
                f"{tbl_name}: "
                f"{row_count} rows"
            )


# ============================================================
# 25. COMPLETE
# ============================================================

print(
    "Migration completed successfully."
)