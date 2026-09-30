"""Build the CESS CSV from Attachment 9a/9f, then optionally load the reviewed CSV.

Usage:

    python src/cess_contribution.py

        Excel -> output/cess_combined.csv


    python src/cess_contribution.py --load-db

        Existing reviewed CSV -> PostgreSQL; no Excel read


Rules:
- field_id must already exist in am_field.
- Missing/unregistered fields remain NULL.
- Never generate a replacement field_id.
- The staging row id is omitted from INSERT so the existing DB default handles it.
"""

import argparse
import csv
import getpass
import os
from pathlib import Path
from uuid import UUID

import psycopg
import yaml
from openpyxl import load_workbook

from src.field_mapping import FIELD_IDS


# ============================================================
# PATHS / CONFIG
# ============================================================

BASE = Path(__file__).resolve().parent.parent

CONFIG = yaml.safe_load(
    (BASE / "environment.yaml").read_text(encoding="utf-8")
) or {}

DB = CONFIG["database"]

SOURCE = BASE / CONFIG["source"]["workbook"]

OUTPUT = BASE / CONFIG.get(
    "output", {}
).get(
    "directory",
    "output",
)

CSV_PATH = OUTPUT / "cess_combined.csv"

MAPPING_PATH = BASE / CONFIG.get(
    "mapping", {}
).get(
    "metadata",
    "pac_name_mapping.yaml",
)


# ============================================================
# DATABASE CONFIG
# ============================================================

DB_SCHEMA = DB.get(
    "schema",
    "amir",
)

TABLES = CONFIG.get(
    "tables",
    {},
)

CESS_TABLE = TABLES.get(
    "cess_combined",
    "abd_cess_combined_stg",
)


# ============================================================
# COLUMNS
# ============================================================

PSC_COLUMNS = (
    "psc_contractor",
    "psc_focal_point",
    "region_psc",
    "psc_expiry",
    "awpb_year",
)

OTHER_METADATA = (
    "field_id",
    "sheet_id",
    "submission_date",
    "region",
    "operator",
)

AUDIT_COLUMNS = (
    "id",
    "created_at",
    "created_by",
    "updated_at",
    "updated_by",
    "delete_at",
    "delete_by",
)

MEASURE_COLUMNS = (
    "work_amount_9a",
    "cost_estimate_9a",
    "cumulative_payment_9a",
    "cess_payment_9a",
    "remaining_payment_9a",
    "crude_production_kbd",
    "gas_production_mmscfd",
    "annual_production_moeb",
    "remaining_production_moeb",
    "cost_estimate_9f",
    "cumulative_payment_9f",
    "payment_due_9f",
)

COLUMNS = (
    "field_name",
    "year",
    *MEASURE_COLUMNS,
    *PSC_COLUMNS,
    *OTHER_METADATA,
    *AUDIT_COLUMNS,
)


# ============================================================
# HELPER FUNCTIONS
# ============================================================

def number(value):
    """Convert Excel numeric values to float."""

    if value is None or (
        isinstance(value, str)
        and value.strip() in ("", "-")
    ):
        return None

    return float(value)


def valid_year(value):
    """Return True when the value is a valid year."""

    return (
        isinstance(value, (int, float))
        and not isinstance(value, bool)
        and 2000 <= value <= 2100
    )


def blank_to_none(value):
    """Convert blank strings to None."""

    if value is None:
        return None

    if isinstance(value, str):
        value = value.strip()
        return value if value else None

    return value


def csv_float(value):
    """Convert CSV value to float or None."""

    value = blank_to_none(value)

    return (
        None
        if value is None
        else float(value)
    )


def csv_int(value):
    """Convert CSV value to integer or None."""

    value = blank_to_none(value)

    return (
        None
        if value is None
        else int(float(value))
    )


# ============================================================
# BUILD CESS CSV
# ============================================================

def build_csv():
    """Extract Attachment 9a/9f and create the review CSV."""

    # --------------------------------------------------------
    # Load PSC metadata mapping
    # --------------------------------------------------------

    mapping = yaml.safe_load(
        MAPPING_PATH.read_text(
            encoding="utf-8"
        )
    ) or {}

    metadata = (
        mapping
        .get("workbooks", {})
        .get(SOURCE.name)
    )

    if not metadata or any(
        metadata.get(column) in (None, "")
        for column in PSC_COLUMNS
    ):
        raise ValueError(
            f"Incomplete PSC metadata mapping for {SOURCE.name}"
        )

    # --------------------------------------------------------
    # Load workbook
    # --------------------------------------------------------

    if not SOURCE.exists():
        raise FileNotFoundError(
            f"Workbook not found: {SOURCE}"
        )

    book = load_workbook(
        SOURCE,
        data_only=True,
    )

    a = book["Attachment 9a"]
    f = book["Attachment 9f"]

    # ========================================================
    # ATTACHMENT 9a(i)
    # ========================================================

    summary = {}

    for row in range(5, 14):

        field = str(
            a.cell(row, 2).value or ""
        ).strip()

        if not field:
            continue

        if field in summary:
            print(
                f"Duplicate 9a(i) {field}: "
                f"using row {row}"
            )

        summary[field] = dict(
            zip(
                MEASURE_COLUMNS[1:5],
                [
                    number(
                        a.cell(
                            row,
                            col,
                        ).value
                    )
                    for col in range(
                        3,
                        7,
                    )
                ],
            )
        )

    # ========================================================
    # ATTACHMENT 9a(ii)
    # ========================================================

    work = {}

    for row in range(29, 36):

        field = str(
            a.cell(row, 2).value or ""
        ).strip()

        if not field:
            continue

        for col in range(3, 8):

            year = a.cell(
                28,
                col,
            ).value

            if valid_year(year):

                work[
                    (
                        field,
                        int(year),
                    )
                ] = number(
                    a.cell(
                        row,
                        col,
                    ).value
                )

    # ========================================================
    # ATTACHMENT 9f
    # ========================================================

    headers = [
        row
        for row in range(
            1,
            f.max_row + 1,
        )
        if str(
            f.cell(row, 2).value or ""
        ).strip().upper() == "FIELD"
    ]

    annual = {}

    for i, start in enumerate(headers):

        field = str(
            f.cell(start, 3).value or ""
        ).strip()

        end = (
            headers[i + 1]
            if i + 1 < len(headers)
            else f.max_row + 1
        )

        for row in range(
            start + 1,
            end,
        ):

            year = f.cell(
                row,
                2,
            ).value

            if not valid_year(year):
                continue

            values = [
                number(
                    f.cell(
                        row,
                        col,
                    ).value
                )
                for col in range(
                    3,
                    10,
                )
            ]

            if any(
                value is not None
                for value in values
            ):
                annual[
                    (
                        field,
                        int(year),
                    )
                ] = dict(
                    zip(
                        MEASURE_COLUMNS[5:],
                        values,
                    )
                )

    # ========================================================
    # COMBINE 9a + 9f
    # ========================================================

    keys = sorted(
        set(work)
        | set(annual)
    )

    # Existing expected result
    if len(keys) != 77:
        raise ValueError(
            "Unexpected CESS row count. "
            f"Expected 77, got {len(keys)}."
        )

    OUTPUT.mkdir(
        parents=True,
        exist_ok=True,
    )

    with CSV_PATH.open(
        "w",
        newline="",
        encoding="utf-8-sig",
    ) as handle:

        writer = csv.DictWriter(
            handle,
            fieldnames=COLUMNS,
        )

        writer.writeheader()

        for field, year in keys:

            writer.writerow(
                {
                    # Business key
                    "field_name": field,
                    "year": year,

                    # Existing verified hierarchy mapping
                    "field_id": FIELD_IDS.get(
                        field
                    ),

                    # Attachment 9a(ii)
                    "work_amount_9a": work.get(
                        (
                            field,
                            year,
                        )
                    ),

                    # Attachment 9a(i)
                    **summary.get(
                        field,
                        {},
                    ),

                    # Attachment 9f
                    **annual.get(
                        (
                            field,
                            year,
                        ),
                        {},
                    ),

                    # PSC metadata
                    **{
                        key: metadata[key]
                        for key in PSC_COLUMNS
                    },
                }
            )

    # ========================================================
    # MAPPING REPORT
    # ========================================================

    mapped_fields = sorted(
        {
            field
            for field, _ in keys
            if FIELD_IDS.get(field)
        }
    )

    unmapped_fields = sorted(
        {
            field
            for field, _ in keys
            if not FIELD_IDS.get(field)
        }
    )

    print(
        f"Created {CSV_PATH}: "
        f"{len(keys)} rows"
    )

    print(
        "Verified field_id mapping: "
        + (
            ", ".join(mapped_fields)
            or "none"
        )
    )

    print(
        "Unmapped fields kept NULL: "
        + (
            ", ".join(unmapped_fields)
            or "none"
        )
    )

    print(
        "Created CSV with verified hierarchy "
        "field_id values only; "
        "no missing field_id is generated."
    )


# ============================================================
# DATABASE CONNECTION
# ============================================================

def get_connection():
    """Create PostgreSQL connection."""

    password = DB.get(
        "password"
    )

    if not password:

        env_name = DB.get(
            "password_env",
            "DIC_DB_PASSWORD",
        )

        password = os.environ.get(
            env_name
        )

    if not password:

        password = getpass.getpass(
            "PostgreSQL password: "
        )

    if not password:

        raise RuntimeError(
            "Database password was not provided."
        )

    return psycopg.connect(
        host=DB.get(
            "host",
            "localhost",
        ),
        port=int(
            DB.get(
                "port",
                5432,
            )
        ),
        dbname=DB.get(
            "name",
            "postgres",
        ),
        user=DB.get(
            "user",
            "postgres",
        ),
        password=password,
        connect_timeout=10,
    )


# ============================================================
# READ REVIEWED CSV
# ============================================================

def read_reviewed_csv():
    """Read the existing reviewed CESS CSV."""

    if not CSV_PATH.exists():

        raise FileNotFoundError(
            f"Reviewed CSV not found: "
            f"{CSV_PATH}. "
            "Run python src/cess_contribution.py first."
        )

    with CSV_PATH.open(
        "r",
        newline="",
        encoding="utf-8-sig",
    ) as handle:

        rows = list(
            csv.DictReader(handle)
        )

    if not rows:

        raise ValueError(
            "CESS CSV contains no data rows."
        )

    # --------------------------------------------------------
    # Validate current CSV schema
    # --------------------------------------------------------

    if "field_name" not in rows[0]:

        raise ValueError(
            "CESS CSV does not contain "
            "'field_name'. "
            "Rebuild the CSV using the current "
            "cess_contribution.py before loading."
        )

    if "year" not in rows[0]:

        raise ValueError(
            "CESS CSV does not contain 'year'."
        )

    if len(rows) != 77:

        raise ValueError(
            "Unexpected reviewed CESS row count. "
            f"Expected 77, got {len(rows)}."
        )

    return rows


# ============================================================
# VERIFY FIELD IDs AGAINST am_field
# ============================================================

def verify_registered_field_ids(
    cursor,
    rows,
):
    """
    Verify every non-NULL field_id in the CSV.

    field_id must:
    - be a valid UUID
    - exist in am_field
    - match the same field_name
    """

    pairs = {}

    for row in rows:

        field_name = blank_to_none(
            row.get(
                "field_name"
            )
        )

        field_id = blank_to_none(
            row.get(
                "field_id"
            )
        )

        if not field_name:

            raise ValueError(
                "CESS CSV contains a row "
                "without field_name."
            )

        if not field_id:
            continue

        try:
            UUID(field_id)

        except ValueError as exc:

            raise ValueError(
                f"Invalid field_id UUID for "
                f"{field_name}: {field_id}"
            ) from exc

        existing = pairs.get(
            field_name
        )

        if (
            existing
            and existing != field_id
        ):

            raise ValueError(
                "More than one field_id found "
                f"for {field_name} in CSV."
            )

        pairs[
            field_name
        ] = field_id

    # --------------------------------------------------------
    # Verify against asset hierarchy
    # --------------------------------------------------------

    for (
        field_name,
        field_id,
    ) in pairs.items():

        cursor.execute(
            f"""
            SELECT 1
            FROM {DB_SCHEMA}.am_field
            WHERE id = %s::uuid
              AND TRIM(field_name) = %s
            LIMIT 1
            """,
            (
                field_id,
                field_name,
            ),
        )

        if cursor.fetchone() is None:

            raise ValueError(
                f"field_id for {field_name} is not "
                f"a matching registered record in "
                f"{DB_SCHEMA}.am_field: {field_id}"
            )

    return pairs


# ============================================================
# LOAD REVIEWED CSV INTO DATABASE
# ============================================================

def load_db():
    """Load/upsert reviewed CESS CSV into PostgreSQL."""

    rows = read_reviewed_csv()

    # --------------------------------------------------------
    # Columns written to PostgreSQL
    #
    # id is intentionally omitted.
    # PostgreSQL generates id for new staging rows.
    # Existing id remains unchanged during ON CONFLICT.
    # --------------------------------------------------------

    insert_columns = (
        "field_id",
        "sheet_id",
        "submission_date",
        "region",
        "operator",
        "field_name",
        "year",
        *MEASURE_COLUMNS,
        *PSC_COLUMNS,
    )

    values = []

    for row in rows:

        values.append(
            (
                blank_to_none(
                    row.get(
                        "field_id"
                    )
                ),

                blank_to_none(
                    row.get(
                        "sheet_id"
                    )
                ),

                blank_to_none(
                    row.get(
                        "submission_date"
                    )
                ),

                blank_to_none(
                    row.get(
                        "region"
                    )
                ),

                blank_to_none(
                    row.get(
                        "operator"
                    )
                ),

                blank_to_none(
                    row.get(
                        "field_name"
                    )
                ),

                csv_int(
                    row.get(
                        "year"
                    )
                ),

                *[
                    csv_float(
                        row.get(column)
                    )
                    for column
                    in MEASURE_COLUMNS
                ],

                blank_to_none(
                    row.get(
                        "psc_contractor"
                    )
                ),

                blank_to_none(
                    row.get(
                        "psc_focal_point"
                    )
                ),

                blank_to_none(
                    row.get(
                        "region_psc"
                    )
                ),

                csv_int(
                    row.get(
                        "psc_expiry"
                    )
                ),

                csv_int(
                    row.get(
                        "awpb_year"
                    )
                ),
            )
        )

    # ========================================================
    # DATABASE TRANSACTION
    # ========================================================

    with get_connection() as connection:

        with connection.cursor() as cursor:

            # ------------------------------------------------
            # Verify hierarchy mappings first
            # ------------------------------------------------

            registered = (
                verify_registered_field_ids(
                    cursor,
                    rows,
                )
            )

            # ------------------------------------------------
            # Verify final business-key constraint
            #
            # Expected:
            # UNIQUE(field_name, year)
            # ------------------------------------------------

            cursor.execute(
                f"""
                SELECT 1
                FROM pg_constraint
                WHERE conrelid =
                    '{DB_SCHEMA}.{CESS_TABLE}'::regclass
                  AND contype = 'u'
                  AND pg_get_constraintdef(oid)
                      LIKE '%(field_name, year)%'
                LIMIT 1
                """
            )

            has_unique_constraint = (
                cursor.fetchone()
                is not None
            )

            if not has_unique_constraint:

                raise RuntimeError(
                    f"{DB_SCHEMA}.{CESS_TABLE} "
                    "does not have a "
                    "UNIQUE(field_name, year) "
                    "constraint. No data was loaded. "
                    "Verify the intended business key "
                    "before enabling rerunnable upserts."
                )

            # ------------------------------------------------
            # SQL placeholders
            # ------------------------------------------------

            placeholders = (
                "%s::uuid, "
                "%s::uuid, "
                "%s, %s, %s, %s, %s, "
                + ", ".join(
                    ["%s"]
                    * len(
                        MEASURE_COLUMNS
                    )
                )
                + ", %s, %s, %s, %s, %s"
            )

            # ------------------------------------------------
            # Columns updated when business key already exists
            # ------------------------------------------------

            update_columns = [
                column
                for column
                in insert_columns
                if column
                not in (
                    "field_name",
                    "year",
                )
            ]

            update_sql = ",\n".join(
                f"{column} = EXCLUDED.{column}"
                for column
                in update_columns
            )

            # ------------------------------------------------
            # UPSERT
            # ------------------------------------------------

            sql = f"""
                INSERT INTO {DB_SCHEMA}.{CESS_TABLE} (
                    {", ".join(insert_columns)}
                )
                VALUES ({placeholders})

                ON CONFLICT (
                    field_name,
                    year
                )

                DO UPDATE SET
                    {update_sql},
                    updated_at = CURRENT_TIMESTAMP
            """

            cursor.executemany(
                sql,
                values,
            )

            # ------------------------------------------------
            # Verify table row count
            # ------------------------------------------------

            cursor.execute(
                f"""
                SELECT COUNT(*)
                FROM {DB_SCHEMA}.{CESS_TABLE}
                """
            )

            count = (
                cursor.fetchone()[0]
            )

            # ------------------------------------------------
            # Verify field mapping after load
            # ------------------------------------------------

            cursor.execute(
                f"""
                SELECT
                    field_name,
                    field_id
                FROM {DB_SCHEMA}.{CESS_TABLE}
                GROUP BY
                    field_name,
                    field_id
                ORDER BY
                    field_name
                """
            )

            mappings = (
                cursor.fetchall()
            )

    # ========================================================
    # LOAD REPORT
    # ========================================================

    print(
        f"Verified {len(registered)} "
        "registered field mappings "
        f"against {DB_SCHEMA}.am_field."
    )

    print(
        f"Loaded/upserted {len(rows)} "
        "reviewed CSV rows into "
        f"{DB_SCHEMA}.{CESS_TABLE}."
    )

    print(
        f"Rows currently in table: "
        f"{count}"
    )

    print(
        "Database field_id mapping:"
    )

    for (
        field_name,
        field_id,
    ) in mappings:

        print(
            f"  {field_name}: "
            f"{field_id if field_id else 'NULL'}"
        )

    print(
        "No replacement field_id was "
        "generated for an unregistered field."
    )


# ============================================================
# MAIN
# ============================================================

def main():

    parser = (
        argparse.ArgumentParser()
    )

    parser.add_argument(
        "--load-db",
        action="store_true",
        help=(
            "Load the existing reviewed "
            "CSV into PostgreSQL."
        ),
    )

    args = parser.parse_args()

    if args.load_db:

        load_db()

    else:

        build_csv()


# ============================================================
# ENTRY POINT
# ============================================================

if __name__ == "__main__":
    main()