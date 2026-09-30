from openpyxl import load_workbook

wb = load_workbook(
    "AWPB 2026 - Balingian PSC - Attachment 9 2.xlsx",
    data_only=True,
)

for name, rows in {
    "Attachment 9g": [2, 6],
    "Attachment 9h": [2, 4, 5, 6, 8, 30, 55],
    "Attachment 9b": [2, 4, 5],
}.items():
    ws = wb[name]
    print(f"\n{name}")
    for row in rows:
        print(
            row,
            [(ws.cell(row, col).coordinate, ws.cell(row, col).value)
             for col in range(1, ws.max_column + 1)
             if ws.cell(row, col).value is not None]
        )