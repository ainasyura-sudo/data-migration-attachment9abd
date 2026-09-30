"""Extract Attachment 9 into reviewable CSVs without writing to PostgreSQL."""

import csv
from collections import defaultdict
from pathlib import Path

import yaml
from openpyxl import load_workbook
from src.field_mapping import FIELD_IDS


BASE = Path(__file__).resolve().parent
CONFIG = yaml.safe_load((BASE / "environment.yaml").read_text(encoding="utf-8"))
SOURCE = BASE / CONFIG["source"]["workbook"]
OUTPUT = BASE / CONFIG["output"]["directory"]
MAPPING = yaml.safe_load(
    (BASE / CONFIG["mapping"]["metadata"]).read_text(encoding="utf-8")
)
metadata = MAPPING["workbooks"].get(SOURCE.name)
if not metadata:
    raise ValueError(f"No verified workbook metadata mapping: {SOURCE.name}")
required_metadata = (
    "psc_name", "psc_contractor", "psc_focal_point", "region_psc",
    "psc_expiry", "awpb_year",
)
if any(metadata.get(key) in (None, "") for key in required_metadata):
    raise ValueError("Incomplete workbook PSC metadata mapping")


def number(value):
    if value is None or isinstance(value, str) and value.strip() in ("", "-"):
        return None
    return float(value)


def label(value):
    return str(value).strip() if value is not None else ""


def save(filename, columns, rows):
    OUTPUT.mkdir(parents=True, exist_ok=True)
    with (OUTPUT / filename).open("w", newline="", encoding="utf-8-sig") as stream:
        writer = csv.DictWriter(stream, fieldnames=columns, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)
    print(f"{filename}: {len(rows)} rows")


wb = load_workbook(SOURCE, data_only=True, read_only=True)
sheets = CONFIG["source"]["sheets"]
work_ws = wb[sheets["masterplan_work"]]
activity_ws = wb[sheets["masterplan_activities"]]
summary_ws = wb[sheets["decom_summary"]]
facility_ws = wb[sheets["decom_facility"]]

work = {}
for row in range(29, 36):
    field = label(work_ws.cell(row, 2).value)
    if not field:
        continue
    for col in range(3, 8):
        year = work_ws.cell(28, col).value
        if isinstance(year, (int, float)) and not isinstance(year, bool):
            key = (field, int(year))
            if key in work:
                raise ValueError(f"Duplicate 9a(ii) field/year: {key}")
            work[key] = number(work_ws.cell(row, col).value)

activities = {}
current_field = ""
for row in range(6, activity_ws.max_row + 1):
    field_cell = label(activity_ws.cell(row, 2).value)
    if field_cell:
        current_field = field_cell
    activity = label(activity_ws.cell(row, 5).value)
    if not current_field or not activity:
        continue
    for index, start_col in enumerate((6, 11, 16, 21, 26)):
        year = 2026 + index
        fy = number(activity_ws.cell(row, start_col + 4).value)
        if fy in (None, 0):
            continue
        key = (current_field, year)
        if key in activities:
            raise ValueError(f"Multiple nonzero 9b activities for {key}")
        activities[key] = {
            "source_row_9b": row,
            "asset_9b": label(activity_ws.cell(row, 3).value) or None,
            "activity_9b": activity,
            **{
                f"q{q + 1}_rm_million": number(activity_ws.cell(row, start_col + q).value)
                for q in range(4)
            },
            "activity_fy_rm_million": fy,
            "remarks_9b": label(activity_ws.cell(row, 31).value) or None,
        }

masterplan = [
    {"field": field, "year": year, "work_amount_9a_rm_million": amount,
     "field_id": FIELD_IDS.get(field), "sheet_id": None,
     **activities.get((field, year), {}), **metadata}
    for (field, year), amount in sorted(work.items())
]
unmatched_activities = set(activities) - set(work)
if unmatched_activities:
    raise ValueError(f"9b activities without a 9a(ii) field/year: {sorted(unmatched_activities)}")
master_columns = [
    "field", "field_id", "sheet_id", "year", "work_amount_9a_rm_million", "source_row_9b",
    "asset_9b", "activity_9b", "q1_rm_million", "q2_rm_million",
    "q3_rm_million", "q4_rm_million", "activity_fy_rm_million",
    "remarks_9b", *required_metadata,
]
save("masterplan.csv", master_columns, masterplan)

summary_columns = [
    "field", "platform_rm_million", "well_rm_million",
    "pipeline_rm_million", "floater_rm_million",
    "subsea_facility_rm_million", "terminal_rm_million",
    "total_rt_rm_million", "total_mod_rm_million",
]
summary = []
for row in range(7, 14):
    field = label(summary_ws.cell(row, 2).value)
    if field:
        summary.append(dict(zip(summary_columns, [
            field, *(number(summary_ws.cell(row, col).value) for col in range(3, 11))
        ])))
save("decom_summary.csv", [*summary_columns, *required_metadata],
     [{**row, **metadata} for row in summary])

facility = []
current_field = ""
for row in range(1, facility_ws.max_row + 1):
    header = label(facility_ws.cell(row, 2).value).lower()
    if header == "field":
        current_field = label(facility_ws.cell(row, 3).value)
    elif header == "facility" and current_field:
        facility.append({
            "field": current_field,
            "source_row_9h": row,
            "facility": label(facility_ws.cell(row, 3).value) or None,
            "total_cost_rm_million": number(facility_ws.cell(row, 6).value),
            **metadata,
        })
facility_columns = ["field", "source_row_9h", "facility", "total_cost_rm_million"]
save("decom_facility.csv", [*facility_columns, *required_metadata], facility)

by_field = defaultdict(list)
for row in facility:
    by_field[row["field"]].append(row)
combined = []
for g in summary:
    for h in by_field.get(g["field"]) or [None]:
        combined.append({
            **g,
            "facility": h["facility"] if h else None,
            "facility_cost_9h_rm_million": h["total_cost_rm_million"] if h else None,
            **metadata,
        })
save("decom_combined.csv", [*summary_columns, "facility",
     "facility_cost_9h_rm_million", "field_id", "sheet_id", *required_metadata],
     [{**row, "field_id": FIELD_IDS.get(row["field"]), "sheet_id": None}
      for row in combined])
